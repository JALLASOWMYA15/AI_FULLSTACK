import os
import re
import json
import hashlib
import sqlite3
from pathlib import Path
from datetime import datetime

import streamlit as st
import fitz
import pytesseract

from PIL import Image
from docx import Document

import chromadb
from sentence_transformers import SentenceTransformer
from transformers import pipeline


# ============================================================
# CONFIGURATION
# ============================================================

APP_TITLE = "AI Personal Digital File Organizer"

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
CHROMA_DIR = DATA_DIR / "chroma"
DB_PATH = DATA_DIR / "files.db"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
CHROMA_DIR.mkdir(parents=True, exist_ok=True)


SUPPORTED_EXTENSIONS = [
    "pdf",
    "jpg",
    "jpeg",
    "png",
    "txt",
    "docx"
]

CATEGORIES = [
    "College Notes",
    "Certificates",
    "ID Documents",
    "Assignments",
    "Bills/Receipts",
    "Screenshots",
    "Photos",
    "Projects",
    "Personal Documents",
    "Other"
]

TAG_KEYWORDS = {
    "Python": ["python", "numpy", "pandas", "django", "flask"],
    "Java": ["java", "jdk", "jvm"],
    "C Programming": ["c language", "c programming", "pointer", "structure"],
    "College": ["college", "university", "semester", "subject", "student"],
    "Assignment": ["assignment", "homework", "question", "submission"],
    "Project": ["project", "implementation", "system", "application"],
    "Certificate": ["certificate", "certification", "certified"],
    "Payment": ["payment", "paid", "transaction", "amount"],
    "Receipt": ["receipt", "invoice", "bill", "gst"],
    "Exam": ["exam", "examination", "mid", "semester exam"],
    "AI": ["artificial intelligence", "machine learning", "deep learning", "ai"],
    "Data Science": ["data science", "data analytics", "dataset"],
    "Important": ["important", "urgent"],
    "Internship": ["internship", "intern"],
    "GitHub": ["github", "repository", "git"],
}


# ============================================================
# STREAMLIT CONFIG
# ============================================================

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📁",
    layout="wide"
)


# ============================================================
# DATABASE
# ============================================================

def get_connection():
    return sqlite3.connect(DB_PATH)


def init_database():
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT UNIQUE,
            file_path TEXT,
            file_type TEXT,
            file_size REAL,
            category TEXT,
            tags TEXT,
            title TEXT,
            summary TEXT,
            extracted_text TEXT,
            metadata TEXT,
            upload_date TEXT
        )
    """)

    connection.commit()
    connection.close()


def file_exists(filename):
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute(
        "SELECT id FROM files WHERE filename = ?",
        (filename,)
    )

    result = cursor.fetchone()

    connection.close()

    return result is not None


def insert_file(
    filename,
    file_path,
    file_type,
    file_size,
    category,
    tags,
    title,
    summary,
    extracted_text,
    metadata
):
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        INSERT INTO files (
            filename,
            file_path,
            file_type,
            file_size,
            category,
            tags,
            title,
            summary,
            extracted_text,
            metadata,
            upload_date
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        filename,
        str(file_path),
        file_type,
        file_size,
        category,
        json.dumps(tags),
        title,
        summary,
        extracted_text,
        json.dumps(metadata),
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ))

    connection.commit()
    connection.close()


def get_all_files():
    connection = get_connection()
    connection.row_factory = sqlite3.Row

    cursor = connection.cursor()

    cursor.execute("""
        SELECT * FROM files
        ORDER BY upload_date DESC
    """)

    rows = cursor.fetchall()

    connection.close()

    return rows


def get_file(file_id):
    connection = get_connection()
    connection.row_factory = sqlite3.Row

    cursor = connection.cursor()

    cursor.execute(
        "SELECT * FROM files WHERE id = ?",
        (file_id,)
    )

    row = cursor.fetchone()

    connection.close()

    return row


def update_file(file_id, category, tags):
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute("""
        UPDATE files
        SET category = ?, tags = ?
        WHERE id = ?
    """, (
        category,
        json.dumps(tags),
        file_id
    ))

    connection.commit()
    connection.close()


def delete_file_from_database(file_id):
    connection = get_connection()
    cursor = connection.cursor()

    cursor.execute(
        "DELETE FROM files WHERE id = ?",
        (file_id,)
    )

    connection.commit()
    connection.close()


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf_text(file_path):
    text = ""

    try:
        pdf = fitz.open(file_path)

        for page_number, page in enumerate(pdf):
            page_text = page.get_text()

            if page_text.strip():
                text += f"\n--- Page {page_number + 1} ---\n"
                text += page_text

        pdf.close()

    except Exception as error:
        st.warning(f"PDF extraction error: {error}")

    return text


def extract_image_text(file_path):
    try:
        image = Image.open(file_path)

        text = pytesseract.image_to_string(image)

        return text

    except Exception as error:
        st.warning(f"OCR error: {error}")
        return ""


def extract_txt_text(file_path):
    try:
        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:
            return file.read()

    except Exception as error:
        st.warning(f"Text extraction error: {error}")
        return ""


def extract_docx_text(file_path):
    text_parts = []

    try:
        document = Document(file_path)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:
                    row_text.append(cell.text)

                text_parts.append(" | ".join(row_text))

    except Exception as error:
        st.warning(f"DOCX extraction error: {error}")

    return "\n".join(text_parts)


def extract_text(file_path):

    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    elif extension in [".jpg", ".jpeg", ".png"]:
        return extract_image_text(file_path)

    elif extension == ".txt":
        return extract_txt_text(file_path)

    elif extension == ".docx":
        return extract_docx_text(file_path)

    return ""


# ============================================================
# AI CATEGORY DETECTION
# ============================================================

def classify_text(text, filename):

    content = (
        filename + " " + text
    ).lower()

    scores = {}

    rules = {
        "College Notes": [
            "notes",
            "lecture",
            "chapter",
            "unit",
            "subject",
            "class notes"
        ],

        "Certificates": [
            "certificate",
            "certification",
            "certified",
            "completion certificate"
        ],

        "ID Documents": [
            "aadhaar",
            "passport",
            "identity",
            "id card",
            "driving license",
            "pan card"
        ],

        "Assignments": [
            "assignment",
            "homework",
            "question",
            "submission",
            "lab record"
        ],

        "Bills/Receipts": [
            "receipt",
            "invoice",
            "bill",
            "gst",
            "total amount",
            "payment"
        ],

        "Screenshots": [
            "screenshot",
            "screen shot"
        ],

        "Projects": [
            "project",
            "implementation",
            "software project",
            "final year project"
        ],

        "Personal Documents": [
            "personal",
            "application",
            "letter",
            "address"
        ],

        "Photos": [
            "photo",
            "photograph"
        ]
    }

    for category, keywords in rules.items():

        score = 0

        for keyword in keywords:

            if keyword in content:
                score += 1

        scores[category] = score

    if max(scores.values()) == 0:
        return "Other"

    return max(
        scores,
        key=scores.get
    )


# ============================================================
# AI IMAGE UNDERSTANDING
# ============================================================

@st.cache_resource
def load_image_classifier():

    try:

        classifier = pipeline(
            "zero-shot-image-classification",
            model="openai/clip-vit-base-patch32"
        )

        return classifier

    except Exception as error:

        st.warning(
            f"Image AI model could not be loaded: {error}"
        )

        return None


def classify_image(file_path):

    try:

        classifier = load_image_classifier()

        if classifier is None:
            return "Photos", []

        image = Image.open(file_path).convert("RGB")

        labels = [
            "a screenshot",
            "a scanned document",
            "a certificate",
            "an identity document",
            "a receipt or bill",
            "a college assignment",
            "college notes",
            "a project document",
            "a personal photograph",
            "a general photo"
        ]

        results = classifier(
            image,
            candidate_labels=labels
        )

        best = results[0]["label"]

        mapping = {
            "a screenshot": "Screenshots",
            "a scanned document": "Personal Documents",
            "a certificate": "Certificates",
            "an identity document": "ID Documents",
            "a receipt or bill": "Bills/Receipts",
            "a college assignment": "Assignments",
            "college notes": "College Notes",
            "a project document": "Projects",
            "a personal photograph": "Photos",
            "a general photo": "Photos"
        }

        category = mapping.get(
            best,
            "Photos"
        )

        image_tags = [
            result["label"]
            for result in results[:3]
        ]

        return category, image_tags

    except Exception as error:

        st.warning(
            f"Image classification failed: {error}"
        )

        return "Photos", []


# ============================================================
# TAG GENERATION
# ============================================================

def generate_tags(text, filename):

    content = (
        filename + " " + text
    ).lower()

    tags = []

    for tag, keywords in TAG_KEYWORDS.items():

        for keyword in keywords:

            if keyword in content:

                if tag not in tags:
                    tags.append(tag)

                break

    return tags[:10]


# ============================================================
# DATE EXTRACTION
# ============================================================

def extract_dates(text):

    patterns = [
        r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b",
        r"\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b",
        r"\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{2,4}\b"
    ]

    dates = []

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text,
            flags=re.IGNORECASE
        )

        dates.extend(matches)

    return list(dict.fromkeys(dates))[:10]


# ============================================================
# TITLE + SUMMARY
# ============================================================

def generate_title(filename, text):

    clean_filename = Path(filename).stem

    clean_filename = clean_filename.replace(
        "_",
        " "
    ).replace(
        "-",
        " "
    )

    if len(clean_filename.strip()) > 3:
        return clean_filename.strip().title()

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if lines:
        return lines[0][:100]

    return "Untitled Document"


def generate_summary(text):

    clean_text = " ".join(
        text.split()
    )

    if not clean_text:
        return "No text could be extracted from this file."

    if len(clean_text) <= 500:
        return clean_text

    return clean_text[:500] + "..."


# ============================================================
# COMPLETE FILE ANALYSIS
# ============================================================

def analyze_file(file_path):

    filename = Path(file_path).name
    extension = Path(file_path).suffix.lower()

    extracted_text = extract_text(file_path)

    category = classify_text(
        extracted_text,
        filename
    )

    vision_tags = []

    if extension in [
        ".jpg",
        ".jpeg",
        ".png"
    ]:

        image_category, vision_tags = classify_image(
            file_path
        )

        if not extracted_text.strip():
            category = image_category

        elif category == "Other":
            category = image_category

    tags = generate_tags(
        extracted_text,
        filename
    )

    for tag in vision_tags:

        if tag not in tags:
            tags.append(tag)

    title = generate_title(
        filename,
        extracted_text
    )

    summary = generate_summary(
        extracted_text
    )

    dates = extract_dates(
        extracted_text
    )

    metadata = {
        "filename": filename,
        "extension": extension,
        "dates": dates,
        "vision_tags": vision_tags,
        "processed_at": datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    }

    return {
        "filename": filename,
        "file_type": extension.replace(".", "").upper(),
        "category": category,
        "tags": tags,
        "title": title,
        "summary": summary,
        "extracted_text": extracted_text,
        "metadata": metadata
    }


# ============================================================
# CHROMADB + SEMANTIC SEARCH
# ============================================================

@st.cache_resource
def load_embedding_model():

    return SentenceTransformer(
        "all-MiniLM-L6-v2"
    )


@st.cache_resource
def load_chroma():

    client = chromadb.PersistentClient(
        path=str(CHROMA_DIR)
    )

    collection = client.get_or_create_collection(
        name="personal_files",
        metadata={
            "hnsw:space": "cosine"
        }
    )

    return collection


def create_search_text(
    filename,
    category,
    tags,
    title,
    summary,
    extracted_text
):

    return f"""
    Filename: {filename}
    Category: {category}
    Tags: {' '.join(tags)}
    Title: {title}
    Summary: {summary}
    Content: {extracted_text}
    """


def create_document_id(filename):

    return hashlib.sha256(
        filename.encode("utf-8")
    ).hexdigest()


def index_file(file_id, analysis):

    try:

        model = load_embedding_model()
        collection = load_chroma()

        search_text = create_search_text(
            analysis["filename"],
            analysis["category"],
            analysis["tags"],
            analysis["title"],
            analysis["summary"],
            analysis["extracted_text"]
        )

        embedding = model.encode(
            search_text
        ).tolist()

        document_id = create_document_id(
            analysis["filename"]
        )

        collection.upsert(
            ids=[document_id],
            embeddings=[embedding],
            documents=[search_text],
            metadatas=[{
                "file_id": str(file_id),
                "filename": analysis["filename"],
                "category": analysis["category"],
                "title": analysis["title"]
            }]
        )

        return True

    except Exception as error:

        st.warning(
            f"Search indexing failed: {error}"
        )

        return False


def search_files(query, number=10):

    try:

        model = load_embedding_model()
        collection = load_chroma()

        if collection.count() == 0:
            return []

        query_embedding = model.encode(
            query
        ).tolist()

        results = collection.query(
            query_embeddings=[query_embedding],
            n_results=min(
                number,
                collection.count()
            )
        )

        output = []

        ids = results.get(
            "ids",
            [[]]
        )[0]

        distances = results.get(
            "distances",
            [[]]
        )[0]

        for index, document_id in enumerate(ids):

            file_id = int(
                results["metadatas"][0][index]["file_id"]
            )

            row = get_file(
                file_id
            )

            if row:

                distance = distances[index]

                similarity = max(
                    0,
                    1 - distance
                )

                output.append({
                    "file": row,
                    "similarity": similarity
                })

        return output

    except Exception as error:

        st.error(
            f"Search error: {error}"
        )

        return []


def delete_from_search(filename):

    try:

        collection = load_chroma()

        document_id = create_document_id(
            filename
        )

        collection.delete(
            ids=[document_id]
        )

    except Exception:
        pass


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def format_size(size_bytes):

    if size_bytes < 1024:
        return f"{size_bytes:.0f} B"

    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.2f} KB"

    return f"{size_bytes / (1024 * 1024):.2f} MB"


def display_file_preview(row):

    file_path = Path(
        row["file_path"]
    )

    if not file_path.exists():

        st.error(
            "File is missing from storage."
        )

        return

    extension = file_path.suffix.lower()

    if extension in [
        ".jpg",
        ".jpeg",
        ".png"
    ]:

        st.image(
            str(file_path),
            width=600
        )

    elif extension == ".pdf":

        with open(
            file_path,
            "rb"
        ) as file:

            st.download_button(
                "📥 Download PDF",
                file,
                file_name=row["filename"],
                mime="application/pdf"
            )

    else:

        st.info(
            "Preview is not available for this file type."
        )

    st.subheader("📄 Extracted Text")

    text = row["extracted_text"]

    if text:

        st.text_area(
            "Content",
            text,
            height=300,
            label_visibility="collapsed"
        )

    else:

        st.info(
            "No text was extracted."
        )


def show_file_card(row):

    tags = json.loads(
        row["tags"]
    ) if row["tags"] else []

    with st.container(border=True):

        col1, col2, col3 = st.columns(
            [4, 2, 1]
        )

        with col1:

            st.subheader(
                f"📄 {row['title']}"
            )

            st.write(
                f"**File:** {row['filename']}"
            )

            st.write(
                f"**Category:** {row['category']}"
            )

            if tags:

                st.write(
                    "**Tags:** "
                    + ", ".join(tags)
                )

        with col2:

            st.write(
                f"**Type:** {row['file_type']}"
            )

            st.write(
                f"**Size:** {format_size(row['file_size'])}"
            )

            st.write(
                f"**Uploaded:** {row['upload_date']}"
            )

        with col3:

            if st.button(
                "View",
                key=f"view_{row['id']}"
            ):

                st.session_state[
                    "selected_file"
                ] = row["id"]

                st.rerun()


# ============================================================
# INITIALIZE
# ============================================================

init_database()


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title("📁 Digital File Organizer")

st.sidebar.caption(
    "AI-powered personal document management"
)

page = st.sidebar.radio(
    "Navigation",
    [
        "Dashboard",
        "Upload Files",
        "All Files",
        "Search",
        "Categories",
        "Settings"
    ]
)

st.sidebar.divider()

st.sidebar.info(
    "Upload your files and let AI organize, "
    "analyze and index them automatically."
)


# ============================================================
# DASHBOARD
# ============================================================

if page == "Dashboard":

    st.title("📊 Dashboard")

    files = get_all_files()

    total_files = len(files)

    pdf_count = sum(
        1
        for file in files
        if file["file_type"] == "PDF"
    )

    image_count = sum(
        1
        for file in files
        if file["file_type"] in [
            "JPG",
            "JPEG",
            "PNG"
        ]
    )

    document_count = sum(
        1
        for file in files
        if file["file_type"] in [
            "TXT",
            "DOCX"
        ]
    )

    total_size = sum(
        file["file_size"]
        for file in files
    )

    col1, col2, col3, col4, col5 = st.columns(5)

    col1.metric(
        "Total Files",
        total_files
    )

    col2.metric(
        "PDFs",
        pdf_count
    )

    col3.metric(
        "Images",
        image_count
    )

    col4.metric(
        "Documents",
        document_count
    )

    col5.metric(
        "Storage",
        format_size(total_size)
    )

    st.divider()

    st.subheader("🚀 How It Works")

    workflow = """
    Upload File
        ↓
    Detect File Type
        ↓
    Extract Text / OCR
        ↓
    AI Content Understanding
        ↓
    Generate Category + Tags + Summary
        ↓
    Create Embeddings
        ↓
    Store in ChromaDB
        ↓
    Semantic Search
    """

    st.code(
        workflow,
        language="text"
    )

    st.subheader("🕒 Recent Files")

    if not files:

        st.info(
            "No files uploaded yet."
        )

    else:

        for row in files[:5]:
            show_file_card(row)


# ============================================================
# UPLOAD FILES
# ============================================================

elif page == "Upload Files":

    st.title("📤 Upload Files")

    st.write(
        "Upload PDFs, images, text files or Word documents."
    )

    uploaded_files = st.file_uploader(
        "Choose files",
        type=SUPPORTED_EXTENSIONS,
        accept_multiple_files=True
    )

    if uploaded_files:

        if st.button(
            "🚀 Process Files",
            type="primary"
        ):

            progress = st.progress(0)

            total = len(
                uploaded_files
            )

            successful = 0

            for index, uploaded_file in enumerate(
                uploaded_files
            ):

                filename = Path(
                    uploaded_file.name
                ).name

                st.write(
                    f"### Processing: {filename}"
                )

                if file_exists(filename):

                    st.warning(
                        f"Skipped duplicate file: {filename}"
                    )

                    progress.progress(
                        (index + 1) / total
                    )

                    continue

                save_path = (
                    UPLOAD_DIR / filename
                )

                try:

                    with open(
                        save_path,
                        "wb"
                    ) as file:

                        file.write(
                            uploaded_file.getbuffer()
                        )

                    with st.spinner(
                        "AI is analyzing the file..."
                    ):

                        analysis = analyze_file(
                            save_path
                        )

                    file_size = (
                        save_path.stat().st_size
                    )

                    insert_file(
                        filename=analysis["filename"],
                        file_path=save_path,
                        file_type=analysis["file_type"],
                        file_size=file_size,
                        category=analysis["category"],
                        tags=analysis["tags"],
                        title=analysis["title"],
                        summary=analysis["summary"],
                        extracted_text=analysis["extracted_text"],
                        metadata=analysis["metadata"]
                    )

                    # Get inserted database ID
                    connection = get_connection()
                    cursor = connection.cursor()

                    cursor.execute(
                        "SELECT id FROM files WHERE filename = ?",
                        (filename,)
                    )

                    row = cursor.fetchone()

                    connection.close()

                    if row:

                        index_file(
                            row[0],
                            analysis
                        )

                    st.success(
                        f"Successfully processed {filename}"
                    )

                    col1, col2, col3 = st.columns(3)

                    col1.write(
                        f"**Category:** {analysis['category']}"
                    )

                    col2.write(
                        "**Tags:** "
                        + ", ".join(
                            analysis["tags"]
                        )
                    )

                    col3.write(
                        f"**Title:** {analysis['title']}"
                    )

                    st.info(
                        f"**Summary:** {analysis['summary']}"
                    )

                    successful += 1

                except Exception as error:

                    if save_path.exists():

                        save_path.unlink()

                    st.error(
                        f"Could not process {filename}: {error}"
                    )

                progress.progress(
                    (index + 1) / total
                )

            st.success(
                f"Finished! {successful} file(s) processed."
            )


# ============================================================
# ALL FILES
# ============================================================

elif page == "All Files":

    st.title("📚 All Files")

    files = get_all_files()

    if not files:

        st.info(
            "No files available. Upload files first."
        )

    else:

        search_filter = st.text_input(
            "🔎 Filter files by name, category or tag"
        )

        filtered_files = []

        for row in files:

            tags = json.loads(
                row["tags"]
            ) if row["tags"] else []

            searchable = (
                row["filename"]
                + " "
                + row["category"]
                + " "
                + " ".join(tags)
            ).lower()

            if (
                not search_filter
                or search_filter.lower()
                in searchable
            ):

                filtered_files.append(row)

        st.write(
            f"Showing {len(filtered_files)} file(s)"
        )

        selected_file_id = st.session_state.get(
            "selected_file"
        )

        if selected_file_id:

            row = get_file(
                selected_file_id
            )

            if row:

                st.subheader(
                    f"📄 {row['filename']}"
                )

                display_file_preview(
                    row
                )

                st.divider()

                current_tags = json.loads(
                    row["tags"]
                ) if row["tags"] else []

                new_category = st.selectbox(
                    "Category",
                    CATEGORIES,
                    index=CATEGORIES.index(
                        row["category"]
                    )
                    if row["category"] in CATEGORIES
                    else 0
                )

                new_tags = st.text_input(
                    "Tags",
                    value=", ".join(
                        current_tags
                    )
                )

                col1, col2 = st.columns(2)

                with col1:

                    if st.button(
                        "💾 Save Changes"
                    ):

                        tag_list = [
                            tag.strip()
                            for tag in new_tags.split(",")
                            if tag.strip()
                        ]

                        update_file(
                            row["id"],
                            new_category,
                            tag_list
                        )

                        st.success(
                            "File details updated."
                        )

                        st.session_state[
                            "selected_file"
                        ] = None

                        st.rerun()

                with col2:

                    if st.button(
                        "🗑️ Delete File"
                    ):

                        delete_from_search(
                            row["filename"]
                        )

                        file_path = Path(
                            row["file_path"]
                        )

                        if file_path.exists():
                            file_path.unlink()

                        delete_file_from_database(
                            row["id"]
                        )

                        st.success(
                            "File deleted."
                        )

                        st.session_state[
                            "selected_file"
                        ] = None

                        st.rerun()

                if st.button(
                    "⬅️ Back to Files"
                ):

                    st.session_state[
                        "selected_file"
                    ] = None

                    st.rerun()

        else:

            for row in filtered_files:

                show_file_card(row)


# ============================================================
# SEARCH
# ============================================================

elif page == "Search":

    st.title("🔍 Semantic Search")

    st.write(
        "Search your files using natural language."
    )

    query = st.text_input(
        "Ask something about your files",
        placeholder="Example: Show my Python assignment"
    )

    examples = [
        "Show my Python assignment",
        "Find my certificates",
        "Show documents related to college",
        "Find payment receipts",
        "Where is my project report?"
    ]

    st.write("**Try:**")

    for example in examples:

        st.caption(
            f"• {example}"
        )

    if query:

        with st.spinner(
            "Searching your files..."
        ):

            results = search_files(
                query
            )

        if not results:

            st.warning(
                "No relevant files found."
            )

        else:

            st.subheader(
                f"🔎 Results for: {query}"
            )

            for result in results:

                row = result["file"]

                similarity = result[
                    "similarity"
                ]

                with st.container(
                    border=True
                ):

                    st.subheader(
                        f"📄 {row['title']}"
                    )

                    st.write(
                        f"**File:** {row['filename']}"
                    )

                    st.write(
                        f"**Category:** {row['category']}"
                    )

                    st.progress(
                        min(
                            max(
                                similarity,
                                0
                            ),
                            1
                        )
                    )

                    st.caption(
                        f"Semantic similarity: "
                        f"{similarity:.2%}"
                    )

                    st.write(
                        row["summary"]
                    )

                    if st.button(
                        "View File",
                        key=f"search_view_{row['id']}"
                    ):

                        st.session_state[
                            "selected_file"
                        ] = row["id"]

                        st.info(
                            "Open 'All Files' to view the selected file."
                        )


# ============================================================
# CATEGORIES
# ============================================================

elif page == "Categories":

    st.title("🗂️ Categories")

    files = get_all_files()

    category_counts = {}

    for row in files:

        category = row["category"]

        category_counts[category] = (
            category_counts.get(
                category,
                0
            ) + 1
        )

    if not category_counts:

        st.info(
            "No categorized files yet."
        )

    else:

        columns = st.columns(3)

        for index, category in enumerate(
            sorted(category_counts)
        ):

            with columns[index % 3]:

                st.metric(
                    category,
                    category_counts[category]
                )


# ============================================================
# SETTINGS
# ============================================================

elif page == "Settings":

    st.title("⚙️ Settings")

    st.subheader(
        "📁 Storage"
    )

    st.write(
        f"Upload directory: `{UPLOAD_DIR}`"
    )

    st.write(
        f"Database: `{DB_PATH}`"
    )

    st.write(
        f"Vector database: `{CHROMA_DIR}`"
    )

    st.divider()

    st.subheader(
        "🤖 AI Components"
    )

    st.write(
        "• PyMuPDF - PDF text extraction"
    )

    st.write(
        "• Tesseract OCR - image text extraction"
    )

    st.write(
        "• Hugging Face Transformers - image understanding"
    )

    st.write(
        "• Sentence Transformers - semantic embeddings"
    )

    st.write(
        "• ChromaDB - vector search"
    )

    st.divider()

    st.subheader(
        "🔐 Privacy"
    )

    st.info(
        "Files are stored locally on your computer. "
        "The application does not require uploading your "
        "documents to an external storage service."
    )

    st.warning(
        "AI models may download from Hugging Face the "
        "first time they are used."
    )